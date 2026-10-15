#!/usr/bin/env python3
"""Revisão automática (1ª passada) das tabelas transcritas manualmente vs fontes primárias.
Trechos literais são SEMPRE extraídos por pdftotext das páginas indicadas; nada é digitado à mão.
Se um trecho não é localizado pelo script, a linha fica 'nao_localizado'."""
import json, re, subprocess, hashlib, csv, os, unicodedata
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

F = '/workspace/revisao/fontes'
OUT = '/workspace/revisao/saida'
AT = '/workspace/doencas_atomizado'
S3 = f'{F}/Sepsis3_Singer2016_JAMA_VoR_via_AmsterdamUMC.pdf'
N1 = f'{F}/NEWS2_RCP2017_Chart1_scoring_system.pdf'
N2 = f'{F}/NEWS2_RCP2017_Chart2_thresholds_triggers.pdf'
DB = f'{F}/DBHA2020_ArqBrasCardiol_116-3_516-658.pdf'
SHA = {p: hashlib.sha256(open(p, 'rb').read()).hexdigest() for p in (S3, N1, N2, DB)}

def pt(pdf, page, layout=False, crop=None):
    cmd = ['pdftotext', '-f', str(page), '-l', str(page)]
    if layout: cmd.append('-layout')
    if crop:
        x, y, w, h = crop; cmd += ['-x', str(x), '-y', str(y), '-W', str(w), '-H', str(h)]
    return subprocess.run(cmd + [pdf, '-'], capture_output=True, text=True, check=True).stdout

def words(pdf, page):
    h = subprocess.run(['pdftotext', '-bbox', '-f', str(page), '-l', str(page), pdf, '-'],
                       capture_output=True, text=True, check=True).stdout
    out = []
    for a, b, c, d, t in re.findall(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">([^<]*)</word>', h):
        t = t.replace('&lt;', '<').replace('&gt;', '>').replace('&amp;', '&')
        out.append((float(a), float(b), float(c), float(d), t))
    return out

def cell(ws, x0, x1, y0, y1, mode='center'):
    sel = [w for w in ws if y0 <= w[1] <= y1 and (x0 <= ((w[0] + w[2]) / 2 if mode == 'center' else w[0]) < x1)]
    sel.sort(key=lambda w: (round(w[1]), w[0]))
    return ' '.join(w[4] for w in sel).strip()

def flat(s): return re.sub(r'\s+', ' ', s).strip()

def find(text, pattern):
    m = re.search(pattern, flat(text), re.I)
    return m.group(0) if m else None

def norm(s):
    s = s.replace('ⱖ', '>=').replace('≥', '>=').replace('ⱕ', '<=').replace('≤', '<=').replace('–', '-').replace('—', '-')
    s = re.sub(r'(?<=\d),(?=\d)', '.', s)
    s = re.sub(r'(?i)\b(sp)?o2\b', 'oxig', s)          # O2/SpO2 não são números
    s = re.sub(r'(?i)\bzero\b', '0', s)
    s = re.sub(r'(?i)\b(\d+) pontos ou mais', r'>=\1', s)
    s = re.sub(r'(?i)\bpresumido (\d+)', r'\1', s)
    return s

TOK = re.compile(r'(>=|<=|<|>)?\s*(\d+(?:\.\d+)?)(?:\s*-\s*(\d+(?:\.\d+)?))?')
def toks(s):
    return [(op or '') + a + ('-' + b if b else '') for op, a, b in TOK.findall(norm(s))]

def split_par(s):
    """separa texto principal e conteúdo entre () ou []"""
    s = norm(s)
    inner = ' '.join(re.findall(r'[\(\[]([^\)\]]*)[\)\]]', s))
    main = re.sub(r'[\(\[][^\)\]]*[\)\]]', ' ', s)
    return main, inner

def nums_only(ts): return [re.sub(r'^(>=|<=|<|>)', '', t) for t in ts]

KW = {  # termo pt na transcrição -> termo en que deve estar no trecho da fonte
    'suporte ventilat': 'respiratory support', 'dopamina': 'dopamine', 'dobutamina': 'dobutamine',
    'qualquer dose': 'any dose', 'norepinefrina': 'norepinephrine', 'epinefrina': 'epinephrine',
    'pam': 'map', 'diurese': 'urine output', 'com o2': 'on oxygen', 'em ar ambiente': 'on air',
    'oxigênio': 'oxygen', 'ar ambiente': 'air', 'alerta': 'alert', 'confusão nova, v, p ou u': 'cvpu',
    'frequência respiratória': 'respiratory rate', 'estado mental': 'altered mentation',
    'glasgow': 'glasgow coma scale', 'sistólica': 'systolic blood pressure',
}
def kw_check(trans, src):
    t, s = trans.lower(), src.lower()
    miss = []
    for k, v in KW.items():
        if k == 'epinefrina' and 'norepinefrina' in t and t.count('epinefrina') == t.count('norepinefrina'):
            continue
        if k == 'ar ambiente' and 'em ar ambiente' in t: continue
        if re.search(r'(?<![a-zà-ú])' + re.escape(k) + r'(?![a-zà-ú])', t) and v not in s:
            miss.append(f'{k}->{v}')
    return miss

def cmp_values(trans, src, allow_missing_inner=True, label=''):
    tm, ti = split_par(trans); sm, si = split_par(src)
    obs = []
    a, b = toks(tm), toks(sm)
    if a != b:
        return 'diverge', f'valores/limiares diferentes: transcrito {a} x fonte {b}'
    if toks(ti):
        if nums_only(toks(ti)) != nums_only(toks(si)):
            return 'diverge', f'valores entre colchetes/parênteses diferentes: transcrito {toks(ti)} x fonte {toks(si)}'
    elif toks(si) and allow_missing_inner:
        obs.append(f'fonte traz também {toks(si)} entre parênteses (unidade alternativa) não transcrito')
    miss = kw_check(trans, label + ' ' + src)
    if miss:
        return 'diverge', 'termo(s) sem correspondente no trecho: ' + '; '.join(miss)
    return 'confere', '; '.join(obs)

def atoms(path):
    return {json.loads(l)['atom_id']: json.loads(l) for l in open(path)}

def field(atom, name):
    m = re.search(r'^' + re.escape(name) + r': (.*)$', atom['descricao_literal'], re.M)
    return m.group(1).strip() if m else None

ROWS = {k: [] for k in ['SOFA', 'qSOFA', 'definicoes_sepsis3', 'NEWS2_pontos', 'NEWS2_risco', 'PA_DBHA2020']}
def add(sheet, atom_id, campo, valor, trecho, pagina, pdf, status, obs=''):
    if status != 'nao_localizado' and not trecho:
        status, obs = 'nao_localizado', (obs + '; ' if obs else '') + 'trecho não localizado automaticamente'
    if status == 'nao_localizado':
        trecho = trecho or ''; pagina = pagina if trecho else ''
    ROWS[sheet].append(dict(atom_id=atom_id, campo=campo, valor_transcrito=valor, trecho_literal_fonte=trecho,
                            pagina=pagina, arquivo_fonte=os.path.basename(pdf), sha256_fonte=SHA[pdf],
                            status_auto=status, observacao=obs))

# ---------------- SOFA (Sepsis-3 Tabela 1, PDF p.5 = JAMA p.804) ----------------
w5 = words(S3, 5)
PG5 = 'PDF p.5 / JAMA 315(8) p.804 (Table 1)'
COLX = [148, 226, 303.9, 381.7, 459.5, 560]
ROWY = {'respiratorio': (118, 135, 'Respiration PaO2/FIO2, mm Hg (kPa)'),
        'coagulacao': (148, 156, 'Coagulation Platelets, ×103/μL'),
        'hepatico': (172, 186, 'Liver Bilirubin, mg/dL (μmol/L)'),
        'cardiovascular': (190, 215, 'Cardiovascular'),
        'sistema-nervoso-central': (228, 240, 'Central nervous system Glasgow Coma Scale score'),
        'renal': (258, 284, 'Renal Creatinine, mg/dL (μmol/L) / Urine output, mL/d')}
# confere que os rótulos das linhas existem na página (extração literal)
labels5 = flat(pt(S3, 5, crop=(60, 95, 85, 195)))
src_cells = {}
for sist, (y0, y1, lab) in ROWY.items():
    for p in range(5):
        c = cell(w5, COLX[p] - 3, COLX[p + 1] - 3, y0, y1, mode='left')
        c = re.sub(r'(?<=[\)\d])\s?[bc]$', '', c).replace(' b ', ' ')  # remove marcas de nota de rodapé b/c
        src_cells[(sist, p)] = c
sofa = {k: v for k, v in atoms(f'{AT}/sepse/criterio_escore_transcrito_manualmente.jsonl').items() if ':sofa:' in k}
seen = set()
for aid, a in sorted(sofa.items()):
    _, _, sist, p = aid.split(':'); p = int(p)
    val = field(a, 'criterio')
    src = src_cells.get((sist, p), '')
    trecho = f'[{ROWY[sist][2]} | Score {p}] {src}' if src else ''
    st, ob = cmp_values(val, src, label=ROWY[sist][2]) if src else ('nao_localizado', 'célula não encontrada')
    if 'Catecholamine' not in ob and sist == 'cardiovascular' and p >= 2:
        foot = find(pt(S3, 5), r'Catecholamine doses are given as μg/kg/min for at least 1 hour\.')
        ob = (ob + '; ' if ob else '') + (f'nota b da fonte: "{foot}"' if foot else 'nota b não localizada')
    add('SOFA', aid, f'{sist} / {p} ponto(s): criterio', val, trecho, PG5, S3, st, ob)
    seen.add((sist, p))
for k in src_cells:
    if k not in seen and src_cells[k]:
        add('SOFA', '(sem átomo)', f'{k[0]} / {k[1]}', '', f'[{ROWY[k[0]][2]} | Score {k[1]}] {src_cells[k]}', PG5, S3,
            'diverge', 'célula da fonte sem transcrição correspondente')
# átomo escore SOFA
esc = atoms(f'{AT}/sepse/escore_transcrito_manualmente.jsonl')
d_sofa = field(esc['sepse:escore:sofa'], 'descricao')
sys_found = [s for s in ['Respiration', 'Coagulation', 'Liver', 'Cardiovascular', 'Central nervous system', 'Renal'] if s in labels5]
add('SOFA', 'sepse:escore:sofa', 'descricao: "Soma de 6 sistemas"', d_sofa,
    'Sistemas listados na Table 1: ' + ' | '.join(sys_found), PG5, S3,
    'confere' if len(sys_found) == 6 else 'diverge', 'contagem de linhas de sistema da Table 1 (a fonte não diz "soma" literalmente; confirmar)')
hdr = cell(w5, 140, 470, 96, 102)
add('SOFA', 'sepse:escore:sofa', 'descricao: "0 a 4 pontos cada"', d_sofa, f'Cabeçalho da Table 1: "Score {hdr}"', PG5, S3,
    'confere' if toks(hdr) == ['0', '1', '2', '3', '4'] else 'diverge', '')
add('SOFA', 'sepse:escore:sofa', 'descricao: "(total 0-24)"', d_sofa, '', '', S3, 'nao_localizado',
    'total 0-24 é derivado aritmeticamente (6x4); não há trecho literal na fonte usada')
sof_src = find(pt(S3, 5), r'Adapted from Vincent et al\.\s?27')
add('SOFA', 'sepse:escore:sofa', 'fonte: Vincent 1996', 'Vincent JL, et al. Intensive Care Med. 1996;22(7):707-710',
    f'Nota a da Table 1: "{sof_src}"' if sof_src else '', PG5, S3, 'nao_localizado',
    'Vincent 1996 (Springer) não acessível publicamente (closed access; download retornou página de desafio HTML). '
    'SOFA conferido contra a Table 1 do Sepsis-3, que declara ser adaptada de Vincent et al. Original NÃO conferido.')

# ---------------- qSOFA (Box 4, PDF p.6 = JAMA p.805) ----------------
PG6 = 'PDF p.6 / JAMA 315(8) p.805'
p6r = pt(S3, 6, layout=True, crop=(300, 60, 295, 80))  # Box 4 (coluna direita, topo)
box4 = [flat(l) for l in p6r.splitlines() if flat(l)]
p6full = pt(S3, 6, layout=True)
qs = {k: v for k, v in atoms(f'{AT}/sepse/criterio_escore_transcrito_manualmente.jsonl').items() if ':qsofa:' in k}
qmap = {'frequencia-respiratoria': r'Respiratory rate ⱖ\s?22/min', 'pressao-sistolica': r'Systolic blood pressure ⱕ\s?100 mm Hg',
        'nivel-consciencia': r'Altered mentation'}
box4t = ' '.join(box4)
for aid, a in sorted(qs.items()):
    var = aid.split(':')[2]; val = field(a, 'criterio')
    tr = find(box4t, qmap[var])
    trecho = f'Box 4. qSOFA (Quick SOFA) Criteria: "{tr}"' if tr else ''
    if var == 'nivel-consciencia':
        g = find(p6full, r'altered mentation \(Glasgow Coma Scale score <15\)')
        if g: trecho += f' || texto, mesma página: "{g}"'
        st, ob = ('confere', 'Box 4 diz só "Altered mentation"; o "(Glasgow < 15)" vem do texto corrido da mesma página') if (tr and g and toks(val) == toks(g)) else ('diverge', 'conferir')
    else:
        st, ob = cmp_values(val, tr or '') if tr else ('nao_localizado', '')
    pts = field(a, 'pontos')
    add('qSOFA', aid, f'{var}: criterio (pontos={pts})', val, trecho, PG6, S3, st,
        (ob + '; ' if ob else '') + '"1 ponto por item" não é literal na fonte (ver linha do escore)')
d_q = field(esc['sepse:escore:qsofa'], 'descricao')
p2 = pt(S3, 2)
t1 = find(p2, r'at least 2 of the following clinical criteria that together constitute a new bedside clinical score termed quickSOFA \(qSOFA\)')
t1b = find(p2, r'adult patients with suspected infection can be rapidly identified as being more likely to have poor outcomes typical of sepsis')
add('qSOFA', 'sepse:escore:qsofa', 'descricao: ">= 2 pontos em suspeita de infecção = maior risco de desfecho ruim"', d_q,
    f'"{t1b} if they have {t1}"' if (t1 and t1b) else '', 'PDF p.2 / JAMA p.801 (Abstract)', S3,
    'confere' if (t1 and t1b) else 'nao_localizado', 'equivalência semântica (tradução) - confirmar')
add('qSOFA', 'sepse:escore:qsofa', 'descricao: "1 ponto por item"', d_q, '', '', S3, 'nao_localizado',
    'fonte fala em "at least 2 of the following clinical criteria"; "1 ponto por item" é interpretação, sem trecho literal')
t3 = find(pt(S3, 8), r'Neither qSOFA nor SOFA is intended to be a stand-alone defi-?\s?nition of sepsis\.')
add('qSOFA', 'sepse:escore:qsofa', 'descricao: "não é critério diagnóstico de sepse"', d_q, f'"{t3}"' if t3 else '',
    'PDF p.8 / JAMA p.807', S3, 'confere' if t3 else 'nao_localizado', 'equivalência semântica (tradução) - confirmar')

# ---------------- Definições (Box 3, PDF p.6) ----------------
p6l = flat(pt(S3, 6, crop=(60, 60, 240, 400)))  # Box 3 (coluna esquerda)
defs = atoms(f'{AT}/sepse/definicao_transcrita_manualmente.jsonl')
ELEM = {
 'sepse:definicao:sepse': [
  ('disfunção orgânica com risco de vida', r'Sepsis is defined as life-threatening organ dysfunction'),
  ('resposta desregulada do hospedeiro à infecção', r'caused by a dysregulated host response to infection\.'),
  ('aumento agudo de 2 pontos ou mais no escore SOFA total', r'Organ dysfunction can be identified as an acute change in total SOFA score ⱖ\s?2 points consequent to the infection\.'),
  ('SOFA basal presumido 0 se não houver disfunção prévia conhecida', r'The baseline SOFA score can be assumed to be zero in patients not known to have preexisting organ dysfunction\.')],
 'sepse:definicao:choque-septico': [
  ('Subgrupo da sepse', r'Septic shock is a subset of sepsis'),
  ('necessidade de vasopressor para manter PAM >= 65 mmHg', r'persisting hypotension requiring vasopressors to maintain MAP ⱖ\s?65 mm Hg'),
  ('lactato sérico > 2 mmol/L (18 mg/dL)', r'serum lactate level >\s?2 mmol/L \(18 mg/dL\)'),
  ('apesar de reposição volêmica adequada', r'despite adequate volume resuscitation\.')],
}
for aid, els in ELEM.items():
    val = field(defs[aid], 'definicao')
    for pt_el, rx in els:
        tr = find(p6l, rx)
        if pt_el.lower() not in val.lower():
            st, ob = 'diverge', 'elemento esperado não está no texto transcrito'
        elif not tr:
            st, ob = 'nao_localizado', ''
        else:
            tv, sv = toks(pt_el), toks(tr)
            st = 'confere' if tv == sv else 'diverge'
            ob = 'tradução pt-BR; números conferem; confirmar equivalência semântica' if st == 'confere' else f'números: {tv} x {sv}'
            if 'vasopressor' in pt_el:
                st, ob = 'diverge', ('OMISSÃO: a fonte diz "persisting hypotension requiring vasopressors"; a transcrição não menciona '
                                     '"hipotensão persistente". Números conferem (PAM >= 65). Provável sem impacto prático; decisão humana.')
        add('definicoes_sepsis3', aid, f'elemento: {pt_el}', val, f'Box 3: "{tr}"' if tr else '', 'PDF p.6 / JAMA 315(8) p.805 (Box 3)', S3, st, ob)

# ---------------- NEWS2 pontos (Chart 1) ----------------
wn = words(N1, 1)
PGN = 'p.1 (Chart 1: The NEWS scoring system)'
CX = [210.4, 261.5, 312.5, 363.5, 414.5, 465.5, 516.5]; CS = [3, 2, 1, 0, 1, 2, 3]
NROW = {'frequencia-respiratoria-irpm': (148, 172, 'Respiration rate (per minute)'),
        'spo2-escala-1': (180, 200, 'SpO2 Scale 1 (%)'),
        'spo2-escala-2-hipercapnia': (205, 240, 'SpO2 Scale 2 (%)'),
        'ar-ambiente-ou-oxigenio': (250, 262, 'Air or oxygen?'),
        'pressao-arterial-sistolica-mmhg': (272, 300, 'Systolic blood pressure (mmHg)'),
        'frequencia-cardiaca-bpm': (305, 325, 'Pulse (per minute)'),
        'consciencia-acvpu': (345, 356, 'Consciousness'),
        'temperatura-c': (365, 390, 'Temperature (°C)')}
lab_text = flat(cell(wn, 100, 190, 140, 400))
ncell = {}
for prm, (y0, y1, lab) in NROW.items():
    for i in range(7):
        lo = (CX[i - 1] + CX[i]) / 2 if i else 190; hi = (CX[i] + CX[i + 1]) / 2 if i < 6 else 560
        ncell[(prm, i)] = cell(wn, lo, hi, y0, y1)
used = set()
for aid, a in sorted(atoms(f'{AT}/sinais_vitais/pontuacao_news2_transcrita_manualmente.jsonl').items()):
    prm = aid.split(':')[2]; faixa = field(a, 'faixa'); pts = int(field(a, 'pontos'))
    if prm not in NROW:
        add('NEWS2_pontos', aid, 'faixa/pontos', f'{faixa} => {pts}', '', '', N1, 'nao_localizado', 'parâmetro não mapeado'); continue
    lab = NROW[prm][2]
    cands = [i for i in range(7) if CS[i] == pts and ncell[(prm, i)]]
    hit = None
    for i in cands:
        if cmp_values(faixa, ncell[(prm, i)])[0] == 'confere': hit = i; break
    if hit is not None:
        used.add((prm, hit))
        add('NEWS2_pontos', aid, f'{lab}: faixa => pontos', f'{faixa} => {pts}', f'[{lab} | Score {CS[hit]}] {ncell[(prm, hit)]}', PGN, N1, 'confere', '')
        continue
    other = [i for i in range(7) if ncell[(prm, i)] and cmp_values(faixa, ncell[(prm, i)])[0] == 'confere']
    if other:
        i = other[0]; used.add((prm, i))
        add('NEWS2_pontos', aid, f'{lab}: faixa => pontos', f'{faixa} => {pts}', f'[{lab} | Score {CS[i]}] {ncell[(prm, i)]}', PGN, N1,
            'diverge', f'faixa encontrada na coluna de {CS[i]} ponto(s), transcrito {pts}')
    else:
        ex = ' || '.join(f'[Score {CS[i]}] {ncell[(prm, i)]}' for i in cands)
        add('NEWS2_pontos', aid, f'{lab}: faixa => pontos', f'{faixa} => {pts}', f'[{lab}] {ex}' if ex else '', PGN, N1,
            'diverge' if ex else 'nao_localizado', 'faixa transcrita não corresponde a nenhuma célula da linha')
for k, v in ncell.items():
    if v and k not in used:
        add('NEWS2_pontos', '(sem átomo)', f'{NROW[k[0]][2]} / Score {CS[k[1]]}', '', f'[{NROW[k[0]][2]} | Score {CS[k[1]]}] {v}', PGN, N1,
            'diverge', 'célula da fonte sem transcrição correspondente')
# nota de consciência (C = confusão nova) — trecho da página oficial RCP não é PDF; registrar só o que está no chart

# ---------------- NEWS2 risco (Chart 2) ----------------
c2 = [flat(l) for l in pt(N2, 1, layout=True).splitlines() if flat(l)]
c2t = ' '.join(c2)
PGR = 'p.1 (Chart 2: NEWS thresholds and triggers)'
RMAP = {'0-4': (r'Aggregate score 0–4 Low', 'baixo', 'low'),
        '3-em-um-unico-parametro': (r'Red score Low–medium Urgent ward-based response\* Score of 3 in any individual parameter', 'baixo-médio', 'low–medium'),
        '5-6': (r'Aggregate score 5–6 Medium', 'médio', 'medium'),
        'ge-7': (r'Aggregate score 7 or more High', 'alto', 'high')}
for aid, a in sorted(atoms(f'{AT}/sinais_vitais/risco_news2_transcrito_manualmente.jsonl').items()):
    key = aid.split(':')[2]; pont = field(a, 'pontuacao'); risco = field(a, 'risco')
    rx, rpt, ren = RMAP[key]
    tr = find(c2t, rx)
    if not tr:
        add('NEWS2_risco', aid, 'pontuacao => risco', f'{pont} => {risco}', '', '', N2, 'nao_localizado', ''); continue
    srcnum = toks(tr.replace('7 or more', '>=7'))
    ok_num = (toks(pont) == srcnum) or (key == '3-em-um-unico-parametro' and toks(pont) == ['3'] and 'Score of 3' in tr)
    ok_r = risco == rpt and ren in tr.lower()
    st = 'confere' if ok_num and ok_r else 'diverge'
    ob = 'coluna "Response" da fonte não foi transcrita' + ('' if st == 'confere' else f'; números {toks(pont)} x {srcnum} / risco {risco} x {ren}')
    if key == '3-em-um-unico-parametro': ob += '; fonte chama esta faixa de "Red score"'
    add('NEWS2_risco', aid, 'pontuacao => risco', f'{pont} => {risco}', f'"{tr}"', PGR, N2, st, ob)

# ---------------- PA DBHA 2020 (Quadro 3.4, PDF p.30 = p.545) ----------------
PGD = 'PDF p.30 / Arq Bras Cardiol 2021;116(3) p.545 (Quadro 3.4)'
d30 = pt(DB, 30, layout=True).splitlines()
title = next((flat(l) for l in d30 if 'Quadro 3.4' in l), '')
hdrl = next((flat(l) for l in d30 if 'PAS (' in l and 'PAD (' in l), '')
LAB = {'pa-otima': 'PA ótima', 'pa-normal': 'PA normal', 'pre-hipertensao': 'Pré-hipertensão',
       'ha-estagio-1': 'HA Estágio 1', 'ha-estagio-2': 'HA Estágio 2', 'ha-estagio-3': 'HA Estágio 3'}
for aid, a in sorted(atoms(f'{AT}/sinais_vitais/classificacao_pa_dbha2020_transcrita_manualmente.jsonl').items()):
    key = aid.split(':')[2]; lab = LAB[key]
    line = next((l for l in d30 if re.match(r'\s*' + re.escape(lab) + r'\s{2,}', l)), None)
    parts = re.split(r'\s{2,}', line.strip()) if line else []
    src = dict(zip(['classificacao', 'pas_mmHg', 'conector', 'pad_mmHg'], parts)) if len(parts) == 4 else {}
    tcls = field(a, 'classificacao')
    for fld in ['classificacao', 'pas_mmHg', 'conector', 'pad_mmHg']:
        tv = field(a, fld); sv = src.get(fld, '')
        trecho = f'linha "{flat(line)}" -> coluna {fld}: "{sv}"' if sv else ''
        if not sv: st, ob = 'nao_localizado', ''
        elif fld == 'classificacao':
            st = 'confere' if tv.lower() == sv.lower() else 'diverge'; ob = '' if tv == sv else 'diferença só de maiúscula/minúscula' if st == 'confere' else ''
        elif fld == 'conector':
            st = 'confere' if tv == sv else 'diverge'; ob = ''
        else:
            st, ob = ('confere', '') if toks(tv) == toks(sv) else ('diverge', f'{toks(tv)} x {toks(sv)}')
        if fld == 'pas_mmHg' and 'mHg' in hdrl and 'PAS (mmHg)' not in hdrl:
            ob = (ob + '; ' if ob else '') + 'cabeçalho da fonte grafa "PAS (mHg)" (erro tipográfico da fonte)'
        add('PA_DBHA2020', aid, fld, tv, trecho, PGD, DB, st, ob)
    if key == 'pa-otima':
        add('PA_DBHA2020', '(todas as 6 linhas)', 'fonte (citação)', field(a, 'fonte'), f'"{title}"', PGD, DB, 'confere',
            'conteúdo (consultório, >=18 anos) confere; a transcrição chama de "Tabela", a fonte é "Quadro 3.4" - sugerir citar o número')
nota = find(pt(DB, 30), r'\*A classificação é definida de acordo com a PA no consultório e pelo nível mais elevado de PA, sistólica ou diastólica\.')
add('PA_DBHA2020', '(todas as 6 linhas)', 'nota de rodapé * do Quadro 3.4', '(não transcrita)', f'"{nota}"' if nota else '', PGD, DB,
    'diverge' if nota else 'nao_localizado', 'OMISSÃO: nota de rodapé * (e ** / *** sobre HA sistólica/diastólica isolada) não foi transcrita; relevante para classificar')

# ---------------- saída ----------------
COLS = ['atom_id', 'campo', 'valor_transcrito', 'trecho_literal_fonte', 'pagina', 'arquivo_fonte', 'sha256_fonte', 'status_auto', 'observacao']
wb = Workbook(); wb.remove(wb.active)
FILL = {'confere': 'C6EFCE', 'diverge': 'FFC7CE', 'nao_localizado': 'FFEB9C'}
resumo = {}
for sh, rows in ROWS.items():
    ws = wb.create_sheet(sh); ws.append(COLS + ['veredito_humano', 'revisor', 'data_revisao'])
    for c in ws[1]: c.font = Font(bold=True)
    for r in rows:
        ws.append([r[c] for c in COLS] + ['', '', ''])
        ws.cell(ws.max_row, 8).fill = PatternFill('solid', fgColor=FILL[r['status_auto']])
    for col, wdt in zip('ABCDEFGHIJKL', [38, 40, 45, 70, 30, 40, 20, 14, 60, 16, 14, 14]):
        ws.column_dimensions[col].width = wdt
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    with open(f'{OUT}/revisao_{sh}.csv', 'w', newline='', encoding='utf-8-sig') as fh:
        wr = csv.DictWriter(fh, fieldnames=COLS); wr.writeheader(); wr.writerows(rows)
    resumo[sh] = {s: sum(r['status_auto'] == s for r in rows) for s in ['confere', 'diverge', 'nao_localizado']}
ws = wb.create_sheet('resumo', 0); ws.append(['tabela', 'confere', 'diverge', 'nao_localizado', 'total'])
for sh, c in resumo.items(): ws.append([sh, c['confere'], c['diverge'], c['nao_localizado'], sum(c.values())])
ws.append([]); ws.append(['arquivo_fonte', 'sha256'])
for p, h in SHA.items(): ws.append([os.path.basename(p), h])
wb.save(f'{OUT}/REVISAO_TABELAS_TRANSCRITAS.xlsx')
json.dump({'resumo': resumo, 'fontes': {os.path.basename(p): h for p, h in SHA.items()}}, open(f'{OUT}/resumo.json', 'w'), ensure_ascii=False, indent=1)
print(json.dumps(resumo, ensure_ascii=False, indent=1))
